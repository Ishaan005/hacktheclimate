"""Audit two EirGrid outage publications and propose reviewed TYTFS switches.

The output describes *scheduled* work, never observed equipment state. Source
workbooks and detailed output belong under ignored ``data/raw`` or ``.cache``.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


ANNUAL_URL = "https://cms.eirgrid.ie/sites/default/files/publications/2026-Transmission-Outage-Programme-20260907.xlsx"
SUMMARY_URL = "https://cms.eirgrid.ie/sites/default/files/publications/Transmission-Outage-Summary-2026-Week-40-41.xlsx"
# HTTP Last-Modified, checked 2026-09-28. These are server file timestamps, not
# independently verified times at which a user could first access the files.
ANNUAL_FILE_TIME = "2026-09-07T13:15:28Z"
SUMMARY_FILE_TIME = "2026-09-17T14:16:45Z"
SUMMARY_START = date(2026, 9, 27)  # first (padding) calendar column in grid


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return datetime.strptime(_text(value), "%d/%m/%Y").date().isoformat()
    except ValueError:
        return None


def read_annual(path: Path) -> dict[str, dict]:
    sheet = load_workbook(path, read_only=True, data_only=True)["GEN_ALL"]
    rows = iter(sheet.values)
    header = next(rows)
    if tuple(header[:8]) != (
        "Outage ID", "Feeder ID", "Calendar Day Duration", "Working Day Duration",
        "Work Days", "Outage Status", "Start", "Finish",
    ):
        raise ValueError("Unexpected annual programme columns")
    result = {}
    for excel_row, cells in enumerate(rows, 2):
        outage_id = _text(cells[0])
        if not outage_id:
            continue
        if outage_id in result:
            raise ValueError(f"Duplicate annual outage ID: {outage_id}")
        result[outage_id] = {
            "outage_id": outage_id,
            "equipment_description": _text(cells[1]),
            "status": _text(cells[5]),
            "start_raw": _text(cells[6]),
            "finish_raw": _text(cells[7]),
            "start_date": _date(cells[6]),
            "finish_date": _date(cells[7]),
            "source_row": excel_row,
        }
    return result


def read_summary(path: Path) -> dict[str, dict]:
    sheet = load_workbook(path, data_only=True).active
    heading = _text(sheet["A1"].value)
    if "28/09/2026 - 11/10/2026" not in heading:
        raise ValueError("Unexpected short-term summary period")
    section = "unknown"
    result = {}
    for row in sheet:
        first = _text(row[0].value)
        if first.startswith("Outages commencing/returning in Week"):
            section = first
        elif first.startswith("Outages spanning Weeks"):
            section = first
        if not first.startswith("TO-"):
            continue
        if first in result:
            raise ValueError(f"Duplicate summary outage ID: {first}")
        days_by_fill: dict[str, list[str]] = {}
        for cell in row[2:18]:
            if cell.fill.patternType != "solid":
                continue
            color = cell.fill.fgColor
            marker = f"{color.type}:{color.indexed if color.type == 'indexed' else color.rgb}"
            day = SUMMARY_START + timedelta(days=cell.column - 3)
            days_by_fill.setdefault(marker, []).append(day.isoformat())
        result[first] = {
            "outage_id": first,
            "plant": _text(row[1].value),
            "section": section,
            "calendar_grid_fills": days_by_fill,
            "status": "listed in short-term outage summary; actual state unconfirmed",
            "source_row": row[0].row,
        }
    return result


def _norm(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def parse_equipment(description: str) -> dict | None:
    """Only parse descriptions with two explicit terminal names and voltages."""
    match = re.fullmatch(
        r"(?P<kind>\d+kV FEEDER|TRANSFORMER|GENERATOR TRANSFORMER|"
        r"DISTRIBUTION TRANSFORMER|TSO CUSTOMER TRANSFORMER) - "
        r"(?P<left>.+?)-(?P<right>.+)-(?P<circuit>[^-]+)",
        description,
        re.IGNORECASE,
    )
    if not match:
        return None
    terminals = []
    for label in (match["left"], match["right"]):
        terminal = re.fullmatch(r"(.+?)\s+(\d{2,3})", label.strip())
        if not terminal:
            return None
        terminals.append((terminal[1], int(terminal[2])))
    return {"kind": match["kind"], "terminals": terminals, "circuit_id": match["circuit"].strip()}


def match_network(
    description: str, buses: Path, branches: Path, transformers: Path,
    reviewed_asset_id: str | None = None,
) -> dict:
    parsed = parse_equipment(description)
    if parsed is None:
        return {"decision": "unresolved", "reason": "equipment endpoints, voltages or circuit not parseable"}
    bus_rows = _read_csv(buses)
    bus_by_id = {row["bus_id"]: row for row in bus_rows}
    candidates = []
    is_feeder = "FEEDER" in parsed["kind"].upper()
    source = _read_csv(branches if is_feeder else transformers)
    for row in source:
        if not is_feeder and _text(row.get("third_bus")) not in ("", "0"):
            continue  # A two-terminal outage cannot safely disable one winding of a 3w unit.
        left = bus_by_id.get(row["from_bus"])
        right = bus_by_id.get(row["to_bus"])
        if left is None or right is None:
            continue
        case_terminals = {
            (_norm(left["name"]), round(float(left["base_kv"]))),
            (_norm(right["name"]), round(float(right["base_kv"]))),
        }
        expected = {(_norm(name), kv) for name, kv in parsed["terminals"]}
        if case_terminals != expected or _norm(row["circuit_id"]) != _norm(parsed["circuit_id"]):
            continue
        candidates.append({
            "asset_type": "branch" if is_feeder else "transformer",
            "asset_id": row.get("asset_id") if is_feeder else row.get("transformer_id"),
            "from_bus": row["from_bus"], "to_bus": row["to_bus"],
            "circuit_id": _norm(row["circuit_id"]),
            "case_in_service": row.get("in_service", ""),
            "transformer_id": row.get("transformer_id") if not is_feeder else None,
        })
    if len(candidates) != 1:
        return {
            "decision": "unresolved",
            "reason": "no exact case match" if not candidates else "multiple exact case matches",
            "candidate_count": len(candidates), "candidates": candidates,
        }
    candidate = candidates[0]
    if reviewed_asset_id and reviewed_asset_id != candidate["asset_id"]:
        return {"decision": "unresolved", "reason": "reviewed asset ID differs from exact candidate", "candidate": candidate}
    result = {
        "decision": "reviewed_scenario_candidate" if reviewed_asset_id else "candidate_requires_manual_review",
        "confidence": "high: unique exact terminals, voltages and circuit; manually reviewed" if reviewed_asset_id else "exact field match; manual review required",
        "candidate": candidates[0],
        "state_warning": "Scenario operation only; neither publication confirms out-of-service state",
    }
    if reviewed_asset_id:
        result["scenario_switch"] = {
            "operation": "set_in_service", "value": False,
            "asset_type": candidate["asset_type"],
            "asset_id": candidate["asset_id"],
            "from_bus": candidate["from_bus"],
            "to_bus": candidate["to_bus"],
            "circuit_id": candidate["circuit_id"],
        }
    return result


def reconcile(annual: dict[str, dict], summary: dict[str, dict]) -> dict:
    shared = annual.keys() & summary.keys()
    missing = summary.keys() - annual.keys()
    # Conflicts are evidence differences, not necessarily errors: a short-term
    # revision may supersede a month-old annual plan.
    conflicts = []
    for outage_id in sorted(shared):
        old, new = annual[outage_id], summary[outage_id]
        marker_dates = set(new["calendar_grid_fills"].get("indexed:23", []))
        start, finish = old["start_date"], old["finish_date"]
        outside = bool(marker_dates and start and finish and any(
            day < start or day > finish for day in marker_dates
        ))
        if outside:
            conflicts.append({
                "outage_id": outage_id,
                "annual_start": start, "annual_finish": finish,
                "summary_marked_days": sorted(marker_dates),
                "annual_source_row": old["source_row"],
                "summary_source_row": new["source_row"],
                "calendar_outside_annual_window": True,
            })
    return {
        "annual_count": len(annual), "summary_count": len(summary),
        "shared_count": len(shared), "summary_missing_from_annual_count": len(missing),
        "annual_absent_from_summary_count": len(annual.keys() - summary.keys()),
        "shared_ids": sorted(shared),
        "summary_missing_from_annual_ids": sorted(missing),
        "summary_missing_from_annual_records": [summary[key] for key in sorted(missing)],
        "shared_annual_statuses": dict(Counter(annual[key]["status"] for key in shared)),
        "date_conflicts": conflicts,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annual", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--buses", type=Path)
    parser.add_argument("--branches", type=Path)
    parser.add_argument("--transformers", type=Path)
    parser.add_argument("--outage-id", help="Inspect one common ID against the TYTFS case")
    parser.add_argument("--reviewed-asset-id", help="Exact TYTFS asset ID confirmed by manual source comparison")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.reviewed_asset_id and not args.outage_id:
        parser.error("--reviewed-asset-id requires --outage-id")
    network_inputs = (args.buses, args.branches, args.transformers)
    if any(network_inputs) and not all(network_inputs):
        parser.error("--buses, --branches and --transformers must be provided together")
    if args.reviewed_asset_id and not all(network_inputs):
        parser.error("--reviewed-asset-id requires all three network CSVs")
    annual, summary = read_annual(args.annual), read_summary(args.summary)
    report = {
        "sources": {
            "annual": {"url": ANNUAL_URL, "http_last_modified": ANNUAL_FILE_TIME},
            "short_term": {"url": SUMMARY_URL, "http_last_modified": SUMMARY_FILE_TIME},
        },
        "publication_caveat": "HTTP Last-Modified is a file timestamp, not verified first availability",
        "historical_cutoff": "Both September 2026 files are excluded from Jan-Aug 2026 backtests",
        "reconciliation": reconcile(annual, summary),
    }
    if args.outage_id:
        old, new = annual.get(args.outage_id), summary.get(args.outage_id)
        if old is None or new is None:
            raise ValueError("Selected outage must occur in both publications")
        report["selected_outage"] = {"annual": old, "short_term": new}
        if all((args.buses, args.branches, args.transformers)):
            report["selected_outage"]["network_match"] = match_network(
                old["equipment_description"], args.buses, args.branches, args.transformers,
                args.reviewed_asset_id,
            )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    main()
