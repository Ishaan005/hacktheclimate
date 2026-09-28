"""Audited ECP generator crosswalk and illustrative network-case allocation.

Outputs derived from EirGrid source files belong under ignored data/raw/.
No name match is promoted to an electrical connection without a review row.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SHEET = "RES & Battery Generation Table"
PROJECT_FIELDS = [
    "source_row", "project_name", "area", "source_node", "generation_type",
    "system_operator", "connection_status", "mec_mw", "candidate_bus_ids",
    "candidate_bus_voltages_kv", "bus_id", "bus_name", "bus_base_kv",
    "review_status", "match_confidence", "allocation_region", "review_evidence",
]


def normalized_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def positive_number(value: object, label: str, *, allow_zero: bool = False) -> float:
    number = float(value)
    if not math.isfinite(number) or number < 0 or (number == 0 and not allow_zero):
        raise ValueError(f"{label} must be finite and {'nonnegative' if allow_zero else 'positive'}")
    return number


def read_buses(path: Path) -> dict[str, dict]:
    buses: dict[str, dict] = {}
    for row in read_csv(path):
        bus_id = str(row["bus_id"]).strip()
        if not bus_id or bus_id in buses:
            raise ValueError(f"empty or repeated bus_id {bus_id!r}")
        buses[bus_id] = {
            "bus_id": bus_id,
            "name": str(row["name"]).strip(),
            "base_kv": positive_number(row["base_kv"], f"bus {bus_id} base_kv"),
            "in_service": str(row.get("in_service", "true")).strip().lower() not in {"false", "0", "no"},
        }
    if not buses:
        raise ValueError("bus table is empty")
    return buses


def read_projects(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET]
        if tuple(sheet.iter_rows(min_row=4, max_row=4, values_only=True))[0][:8] != (
            "Name", "Area", "Node", "Generation", "System", "Connection", "Processing", "Maximum Export"
        ):
            raise ValueError("ECP generation table header changed")
        projects = []
        for cells in sheet.iter_rows(min_row=6, values_only=False):
            row_number = cells[0].row
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            name, area, node, kind, operator, status, _, mec = values[:8]
            if not (name and area and node and kind and operator and status):
                raise ValueError(f"incomplete project row {row_number}")
            status = str(status).strip().lower()
            kind = str(kind).strip().lower()
            operator = str(operator).strip().upper()
            if status not in {"connected", "due to connected"}:
                raise ValueError(f"unexpected connection status at row {row_number}: {status}")
            if kind not in {"wind", "offshore wind", "solar", "battery"} or operator not in {"TSO", "DSO"}:
                raise ValueError(f"unexpected type/operator at row {row_number}: {kind}/{operator}")
            projects.append({
                "source_row": str(row_number), "project_name": str(name).strip(),
                "area": str(area).strip(), "source_node": str(node).strip(),
                "generation_type": kind, "system_operator": operator,
                "connection_status": status,
                "mec_mw": positive_number(mec, f"row {row_number} MEC"),
            })
        return projects
    finally:
        workbook.close()


def build_crosswalk(workbook: Path, buses_csv: Path, reviews_csv: Path | None = None) -> tuple[list[dict], dict]:
    buses = read_buses(buses_csv)
    by_name: dict[str, list[dict]] = defaultdict(list)
    for bus in buses.values():
        by_name[normalized_name(bus["name"])].append(bus)
    source_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()
    reviews: dict[str, dict] = {}
    if reviews_csv:
        for review in read_csv(reviews_csv):
            source_row = review["source_row"].strip()
            if source_row in reviews:
                raise ValueError(f"duplicate review of source row {source_row}")
            if review["workbook_sha256"].strip() != source_hash:
                raise ValueError(f"review row {source_row} was made against another ECP workbook")
            reviews[source_row] = review
    rows = []
    for project in read_projects(workbook):
        matches = sorted(by_name[normalized_name(project["source_node"])], key=lambda bus: bus["bus_id"])
        result = {**project,
                  "candidate_bus_ids": "|".join(bus["bus_id"] for bus in matches),
                  "candidate_bus_voltages_kv": "|".join(str(bus["base_kv"]) for bus in matches),
                  "bus_id": "", "bus_name": "", "bus_base_kv": "",
                  "review_status": "unmatched" if not matches else "ambiguous",
                  "match_confidence": "none" if not matches else "name_only",
                  "allocation_region": project["area"], "review_evidence": ""}
        review = reviews.pop(project["source_row"], None)
        if review:
            if project["connection_status"] != "connected":
                raise ValueError(f"future row {project['source_row']} cannot be accepted")
            bus_id = review["bus_id"].strip()
            status = review["review_status"].strip()
            evidence = review["review_evidence"].strip()
            if status not in {"accepted_proxy", "accepted_verified", "rejected"}:
                raise ValueError(f"invalid review status for row {project['source_row']}")
            if not evidence:
                raise ValueError(f"review row {project['source_row']} needs evidence")
            if status.startswith("accepted"):
                if bus_id not in buses:
                    raise ValueError(f"review row {project['source_row']} uses absent bus {bus_id}")
                bus = buses[bus_id]
                if not bus["in_service"]:
                    raise ValueError(f"review row {project['source_row']} uses out-of-service bus {bus_id}")
                if normalized_name(bus["name"]) != normalized_name(review["bus_name"].strip()):
                    raise ValueError(f"review row {project['source_row']} bus name changed")
                if bus["base_kv"] != float(review["bus_base_kv"]):
                    raise ValueError(f"review row {project['source_row']} bus voltage changed")
                if bus_id not in {match["bus_id"] for match in matches} and not review.get("override_reason", "").strip():
                    raise ValueError(f"review row {project['source_row']} needs override reason for non-name match")
                result.update(bus_id=bus_id, bus_name=bus["name"], bus_base_kv=bus["base_kv"],
                              match_confidence="reviewed_station_proxy" if status == "accepted_proxy" else "reviewed_connection")
                result["allocation_region"] = review["allocation_region"].strip() or project["area"]
            else:
                if bus_id:
                    raise ValueError(f"rejected row {project['source_row']} must have empty bus_id")
            result.update(review_status=status, review_evidence=evidence)
        rows.append(result)
    if reviews:
        raise ValueError(f"reviews reference absent source rows: {sorted(reviews)}")
    summary = {"workbook_sha256": source_hash, "bus_count": len(buses), "project_count": len(rows), "by_status": {}}
    for status in ("connected", "due to connected"):
        subset = [row for row in rows if row["connection_status"] == status]
        summary["by_status"][status] = {}
        for review_status in ("accepted_proxy", "accepted_verified", "ambiguous", "unmatched", "rejected"):
            group = [row for row in subset if row["review_status"] == review_status]
            summary["by_status"][status][review_status] = {
                "rows": len(group), "mec_mw": round(sum(row["mec_mw"] for row in group), 3)
            }
        summary["by_status"][status]["total"] = {
            "rows": len(subset), "mec_mw": round(sum(row["mec_mw"] for row in subset), 3)
        }
        summary["by_status"][status]["by_operator_and_type"] = {
            f"{operator}/{kind}": {
                "rows": len(group), "mec_mw": round(sum(row["mec_mw"] for row in group), 3)
            }
            for operator in ("TSO", "DSO") for kind in ("wind", "offshore wind", "solar", "battery")
            if (group := [row for row in subset if row["system_operator"] == operator and row["generation_type"] == kind])
        }
    return rows, summary


def allocate(crosswalk_csv: Path, renewable_csv: Path, load_weights_csv: Path,
             total_load_mw: float, buses_csv: Path) -> tuple[list[dict], dict]:
    """Allocate explicit regional wind/solar MW and scoped total load to case buses.

    MEC is only a within-group weight. All input MW must be placed or this fails.
    """
    buses = read_buses(buses_csv)
    projects = read_csv(crosswalk_csv)
    accepted = [row for row in projects if row["connection_status"] == "connected"
                and row["review_status"] in {"accepted_proxy", "accepted_verified"}]
    unreviewed_by_group: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in projects:
        if row["connection_status"] == "connected" and row["review_status"] not in {"accepted_proxy", "accepted_verified"}:
            kind = "wind" if row["generation_type"] == "offshore wind" else row["generation_type"]
            unreviewed_by_group[(row["allocation_region"], kind)].append(row)
    by_group: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in accepted:
        if row["bus_id"] not in buses:
            raise ValueError(f"accepted row {row['source_row']} bus is absent")
        if row["generation_type"] in {"wind", "offshore wind", "solar"}:
            kind = "wind" if row["generation_type"] == "offshore wind" else row["generation_type"]
            by_group[(row["allocation_region"], kind)].append(row)
    output: dict[str, dict] = {bus_id: {"bus_id": bus_id, "wind_mw": 0.0, "solar_mw": 0.0, "load_mw": 0.0}
                               for bus_id in buses}
    forecast_totals = {"wind": 0.0, "solar": 0.0}
    seen_groups = set()
    for forecast in read_csv(renewable_csv):
        area = forecast["allocation_region"].strip()
        kind = forecast["generation_type"].strip().lower()
        if kind not in forecast_totals:
            raise ValueError("only wind and solar generation may be allocated; battery dispatch requires a separate assumption")
        key = (area, kind)
        if key in seen_groups:
            raise ValueError(f"repeated forecast group {key}")
        seen_groups.add(key)
        mw = positive_number(forecast["forecast_mw"], f"forecast {key}", allow_zero=True)
        group = by_group[key]
        if mw and not group:
            raise ValueError(f"forecast group {key} has no reviewed connected sites")
        if mw and unreviewed_by_group[key]:
            raise ValueError(f"forecast group {key} includes unreviewed connected sites; define an explicit reviewed subset")
        total_mec = sum(positive_number(row["mec_mw"], f"row {row['source_row']} MEC") for row in group)
        if mw > total_mec + 1e-8:
            raise ValueError(f"forecast group {key} exceeds reviewed MEC {total_mec} MW")
        for row in group:
            output[row["bus_id"]][f"{kind}_mw"] += mw * float(row["mec_mw"]) / total_mec
        forecast_totals[kind] += mw
    weights: dict[str, float] = defaultdict(float)
    load_rows = read_csv(load_weights_csv)
    if not load_rows:
        raise ValueError("load weights table is empty")
    case_load_weights = "p_mw" in load_rows[0]
    for row in load_rows:
        if case_load_weights and str(row.get("in_service", "true")).strip().lower() in {"false", "0", "no"}:
            continue
        bus_id = row["bus_id"].strip()
        if bus_id not in buses or (bus_id in weights and not case_load_weights):
            raise ValueError(f"unknown or repeated load bus {bus_id}")
        value = row["p_mw"] if case_load_weights else row["weight"]
        weights[bus_id] += positive_number(value, f"load weight {bus_id}", allow_zero=True)
    total_weight = sum(weights.values())
    if total_weight <= 0:
        raise ValueError("load weights need a positive sum")
    total_load_mw = positive_number(total_load_mw, "total load MW", allow_zero=True)
    for bus_id, weight in weights.items():
        output[bus_id]["load_mw"] = total_load_mw * weight / total_weight
    rows = [row for row in output.values() if any(row[field] for field in ("wind_mw", "solar_mw", "load_mw"))]
    for row in rows:
        row["net_injection_mw"] = row["wind_mw"] + row["solar_mw"] - row["load_mw"]
    actual = {field: sum(row[field] for row in rows) for field in ("wind_mw", "solar_mw", "load_mw")}
    if not all(math.isclose(actual[f"{kind}_mw"], mw, abs_tol=1e-8) for kind, mw in forecast_totals.items()):
        raise AssertionError("generation allocation failed to reconcile")
    if not math.isclose(actual["load_mw"], total_load_mw, abs_tol=1e-8):
        raise AssertionError("load allocation failed to reconcile")
    return rows, {"input_forecast_mw": forecast_totals, "input_total_load_mw": total_load_mw,
                  "allocated_mw": actual, "net_injection_mw": sum(row["net_injection_mw"] for row in rows),
                  "renewable_method": "MEC shares within explicitly supplied region/type groups; illustrative potential, not measured farm output",
                  "load_method": ("in-service case p_mw by bus" if case_load_weights else "explicit bus weights")
                  + " normalized to total case load assumption",
                  "load_weight_total": total_weight}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    crosswalk = sub.add_parser("crosswalk")
    crosswalk.add_argument("--workbook", type=Path, required=True)
    crosswalk.add_argument("--buses", type=Path, required=True)
    crosswalk.add_argument("--reviews", type=Path)
    crosswalk.add_argument("--output", type=Path, required=True)
    crosswalk.add_argument("--report", type=Path, required=True)
    allocation = sub.add_parser("allocate")
    allocation.add_argument("--crosswalk", type=Path, required=True)
    allocation.add_argument("--buses", type=Path, required=True)
    allocation.add_argument("--renewable-forecast", type=Path, required=True)
    allocation.add_argument("--load-weights", type=Path, required=True)
    allocation.add_argument("--total-load-mw", type=float, required=True)
    allocation.add_argument("--output", type=Path, required=True)
    allocation.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "crosswalk":
        rows, report = build_crosswalk(args.workbook, args.buses, args.reviews)
        write_csv(args.output, PROJECT_FIELDS, rows)
    else:
        rows, report = allocate(args.crosswalk, args.renewable_forecast, args.load_weights,
                                args.total_load_mw, args.buses)
        write_csv(args.output, ["bus_id", "wind_mw", "solar_mw", "load_mw", "net_injection_mw"], rows)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
