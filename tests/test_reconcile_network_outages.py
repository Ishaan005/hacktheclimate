import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from scripts.reconcile_network_outages import (
    match_network, parse_equipment, read_annual, read_summary, reconcile,
)


def _csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def test_exact_network_match_requires_review_and_preserves_ambiguity(tmp_path):
    buses = tmp_path / "buses.csv"
    branches = tmp_path / "branches.csv"
    transformers = tmp_path / "transformers.csv"
    _csv(buses, ["bus_id", "name", "base_kv"], [
        {"bus_id": "1642", "name": "'CASHLA      '", "base_kv": "220.0"},
        {"bus_id": "2522", "name": "'FLAGFORD    '", "base_kv": "220.0"},
    ])
    row = {"asset_id": "1642:2522:1", "from_bus": "1642", "to_bus": "2522",
           "circuit_id": "'1 '", "in_service": "True"}
    fields = list(row)
    _csv(branches, fields, [row])
    _csv(transformers, ["transformer_id", "from_bus", "to_bus", "third_bus", "circuit_id"], [])
    description = "220kV FEEDER - CASHLA 220-FLAGFORD 220-1"

    candidate = match_network(description, buses, branches, transformers)
    assert candidate["decision"] == "candidate_requires_manual_review"
    assert "scenario_switch" not in candidate
    reviewed = match_network(description, buses, branches, transformers, "1642:2522:1")
    assert reviewed["decision"] == "reviewed_scenario_candidate"
    assert reviewed["scenario_switch"] == {
        "operation": "set_in_service", "value": False, "asset_type": "branch",
        "asset_id": "1642:2522:1", "from_bus": "1642", "to_bus": "2522", "circuit_id": "1",
    }
    assert match_network(description, buses, branches, transformers, "wrong")["decision"] == "unresolved"
    _csv(branches, fields, [row, row])
    assert match_network(description, buses, branches, transformers, "1642:2522:1")["reason"] == "multiple exact case matches"
    assert parse_equipment("CAPACITOR - CASHLA 220-CASHLA 220-1") is None


def test_workbook_reconciliation_keeps_source_status_and_grid_markers(tmp_path):
    annual_path = tmp_path / "annual.xlsx"
    summary_path = tmp_path / "summary.xlsx"
    annual = Workbook()
    sheet = annual.active
    sheet.title = "GEN_ALL"
    sheet.append(["Outage ID", "Feeder ID", "Calendar Day Duration", "Working Day Duration",
                  "Work Days", "Outage Status", "Start", "Finish"])
    sheet.append(["TO-1", "220kV FEEDER - CASHLA 220-FLAGFORD 220-1", "2", "2", 2,
                  "Scheduled", "28/09/2026", "29/09/2026"])
    annual.save(annual_path)

    summary = Workbook()
    sheet = summary.active
    sheet["A1"] = "Report covers outages commencing/returning/spanning the period: 28/09/2026 - 11/10/2026"
    sheet["A2"] = "Outages commencing/returning in Week 40"
    sheet["A6"] = "TO-1"
    sheet["B6"] = "CASHLA FLAGFORD"
    sheet["D6"].fill = PatternFill("solid", fgColor="00AAAAAA")
    sheet["E6"].fill = PatternFill("solid", fgColor="00AAAAAA")
    sheet["F6"].fill = PatternFill("solid", fgColor="00AAAAAA")
    sheet["A7"] = "TO-2"
    sheet["B7"] = "UNKNOWN PLANT"
    summary.save(summary_path)

    old, new = read_annual(annual_path), read_summary(summary_path)
    assert old["TO-1"]["status"] == "Scheduled"
    assert old["TO-1"]["start_date"] == "2026-09-28"
    assert new["TO-1"]["status"].endswith("actual state unconfirmed")
    assert new["TO-1"]["calendar_grid_fills"]["rgb:00AAAAAA"] == [
        "2026-09-28", "2026-09-29", "2026-09-30",
    ]
    audit = reconcile(old, new)
    assert (audit["shared_count"], audit["summary_missing_from_annual_count"]) == (1, 1)
    assert audit["summary_missing_from_annual_ids"] == ["TO-2"]
